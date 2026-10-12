# excel_exporter.py
import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _build_workbook(jobs: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws_jobs = wb.active
    ws_jobs.title = "Job Listings"
    ws_jobs.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=15, bold=True, color="1F4E78")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_link = Font(name="Calibri", size=11, color="0563C1", underline="single")
    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_alt = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws_jobs.cell(row=1, column=1, value="Automated Job Search & Match Tracker").font = font_title
    headers = [
        "Job ID",
        "Role Title",
        "Company",
        "Location",
        "Source Board",
        "Posted Date",
        "Match Score",
        "Title Fit",
        "Exp Fit",
        "Semantic Fit",
        "Application Link",
        "Visa Status",
        "Matching Skills",
        "Missing Skills",
    ]
    header_row = 3

    for col_idx, h in enumerate(headers, start=1):
        cell = ws_jobs.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    start_row = 4
    for idx, j in enumerate(jobs, start=start_row):
        ws_jobs.cell(row=idx, column=1, value=f"JOB-{idx - start_row + 101}").alignment = Alignment(
            horizontal="center"
        )
        ws_jobs.cell(row=idx, column=2, value=j.get("title", "N/A"))
        ws_jobs.cell(row=idx, column=3, value=j.get("company", "N/A"))
        ws_jobs.cell(row=idx, column=4, value=j.get("location", "N/A"))
        ws_jobs.cell(row=idx, column=5, value=str(j.get("site", "N/A")).title()).alignment = Alignment(
            horizontal="center"
        )
        ws_jobs.cell(row=idx, column=6, value=str(j.get("date_posted", "N/A"))).alignment = Alignment(
            horizontal="center"
        )

        score_cell = ws_jobs.cell(row=idx, column=7, value=j.get("match_score", 0.70))
        score_cell.number_format = "0%"
        score_cell.alignment = Alignment(horizontal="right")

        for col_offset, key in enumerate(("title_score", "exp_score", "semantic_score"), start=8):
            c = ws_jobs.cell(row=idx, column=col_offset, value=j.get(key, 0))
            c.number_format = "0%"
            c.alignment = Alignment(horizontal="right")

        url = j.get("job_url") or j.get("site_url") or ""
        link_cell = ws_jobs.cell(row=idx, column=11, value="Apply Link" if url else "N/A")
        if url:
            link_cell.hyperlink = url
            link_cell.font = font_link
        link_cell.alignment = Alignment(horizontal="center")

        ws_jobs.cell(row=idx, column=12, value=j.get("visa_status", "N/A")).alignment = Alignment(
            horizontal="center"
        )
        ws_jobs.cell(row=idx, column=13, value=j.get("key_matches", "N/A"))
        ws_jobs.cell(row=idx, column=14, value=j.get("missing_skills", "None"))

        fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
        for c in range(1, len(headers) + 1):
            cell = ws_jobs.cell(row=idx, column=c)
            cell.border = border_thin
            if fill.fill_type:
                cell.fill = fill

    for col in ws_jobs.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_jobs.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws_jobs.column_dimensions["B"].width = 32
    ws_jobs.column_dimensions["M"].width = 28
    ws_jobs.column_dimensions["N"].width = 28
    return wb


def export_jobs_to_excel(jobs: list[dict], file_path: str = "job_search_results.xlsx") -> str:
    wb = _build_workbook(jobs)
    try:
        wb.save(file_path)
        return file_path
    except PermissionError:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_path = f"job_search_results_{timestamp}.xlsx"
        wb.save(fallback_path)
        return fallback_path


def export_jobs_to_excel_bytes(jobs: list[dict]) -> bytes:
    buffer = BytesIO()
    _build_workbook(jobs).save(buffer)
    return buffer.getvalue()
